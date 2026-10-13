"""Rozdziel opisy obwodów i uzgodnij nazwy z lokalną geobazą PRG.

Uruchomienie (Python QGIS z GDAL i openpyxl):
  python rozbij_granice.py
Alternatywnie: Python ArcGIS Pro z aktywną licencją (ArcPy).
"""
from __future__ import annotations

import argparse
import copy
import csv
import difflib
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path


BASE = Path(__file__).resolve().parent
GEO_RESULTS = BASE / 'wyniki_PRG_Geoportal'


def prg_layer_name(names, kind):
    """Jednoznaczny wybór nazw warstw ze starego eksportu i Geoportalu."""
    prefix = 'PRG_PunktyAdresowe' if kind == 'points' else 'PRG_Ulice'
    candidates = [name for name in names if name.startswith(prefix)]
    if len(candidates) != 1:
        raise ValueError(f'Oczekiwano jednej warstwy {prefix}; znaleziono: {candidates}')
    return candidates[0]


def key(text):
    """Normalizacja wyłącznie do porównania; wynik zachowuje pisownię PRG."""
    text = str(text or '').casefold().replace('ł', 'l')
    text = ''.join(c for c in unicodedata.normalize('NFKD', text)
                   if not unicodedata.combining(c))
    text = re.sub(r'\b(?:ksiedza|ks)\b', 'ks', text)
    text = re.sub(r'\b(?:generala|general|gen)\b', 'gen', text)
    text = re.sub(r'\b(?:doktora|dra|dr)\b', 'dr', text)
    text = re.sub(r'\b(?:profesora|prof)\b', 'prof', text)
    text = re.sub(r'\b(?:imienia|im)\b', '', text)
    text = re.sub(r'[^\w\s-]', ' ', text)
    return ' '.join(text.split())


def name_key(text):
    text = key(text)
    # Powtórzone oznaczenia, np. „pl. Plac ...”, „ul. Aleja ...”.
    return re.sub(r'^(?:(?:ul|ulica|al|aleja|aleje|pl|plac|skwer)\s+)+', '', text)


def read_prg(gdb):
    """Odczyt tylko atrybutów, bez modyfikacji geobazy i bez geometrii."""
    fields = ['NAZWA_ULC', 'NAZWA_TER1', 'NAZWA_TER2', 'NAZWA_MSC']
    try:
        from osgeo import ogr
    except ImportError:
        try:
            import arcpy
        except (ImportError, RuntimeError) as exc:
            raise RuntimeError('Uruchom w Pythonie QGIS (GDAL) albo licencjonowanym Pythonie ArcGIS Pro.') from exc
        records = []
        with arcpy.EnvManager(workspace=str(gdb)):
            names = arcpy.ListFeatureClasses() or []
        for layer in [prg_layer_name(names, 'streets'), prg_layer_name(names, 'points')]:
            path = str(gdb / layer)
            available = {f.name for f in arcpy.ListFields(path)}
            chosen = [f for f in fields if f in available]
            with arcpy.da.SearchCursor(path, chosen) as cursor:
                records.extend(dict(zip(chosen, row)) for row in cursor)
        return records
    ogr.UseExceptions()
    dataset = ogr.Open(str(gdb), 0)
    if dataset is None:
        raise ValueError(f'Nie można otworzyć geobazy: {gdb}')
    records = []
    names = [layer.GetName() for layer in dataset]
    for layer_name in [prg_layer_name(names, 'streets'), prg_layer_name(names, 'points')]:
        layer = dataset.GetLayerByName(layer_name)
        if layer is None:
            raise ValueError(f'Brak warstwy {layer_name}')
        available = {f.GetName() for f in layer.schema}
        if 'NAZWA_ULC' not in available:
            raise ValueError(f'Brak NAZWA_ULC w {layer_name}')
        for feature in layer:
            records.append({f: feature.GetField(f) for f in fields if f in available})
    return records


class StreetIndex:
    def __init__(self, records):
        self.literal = defaultdict(set)
        self.exact = defaultdict(set)
        self.alias = defaultdict(set)
        self.names = set()
        for record in records:
            if key(record.get('NAZWA_MSC')) != 'lublin':
                continue
            name = str(record.get('NAZWA_ULC') or '').strip()
            if not name:
                continue
            self.names.add(name)
            self.literal[' '.join(name.casefold().split())].add(name)
            self.exact[key(name)].add(name)
            self.alias[name_key(name)].add(name)
            # PRG przechowuje człony patrona również w oddzielnych polach.
            term1 = str(record.get('NAZWA_TER1') or '').strip()
            term2 = str(record.get('NAZWA_TER2') or '').strip()
            if term1:
                self.alias[name_key(f'{term2} {term1}')].add(name)
            # Brak tytułu w opisie: tylko jednoznaczne dopasowanie.
            without_title = re.sub(r'^(?:dr|ks|prof|gen)\s+', '', name_key(name))
            self.alias[without_title].add(name)
        if not self.names:
            raise ValueError('Nie znaleziono nazw ulic dla Lublina w PRG.')

    def match(self, name):
        matches = (self.literal.get(' '.join(name.casefold().split()))
                   or self.exact.get(key(name)) or self.alias.get(name_key(name), set()))
        return sorted(matches)


def split_description(text):
    """Przecinki wewnątrz nawiasów i listy numerów pozostają przy ulicy."""
    text = re.sub(r'^\s*ULICA\s*:\s*', '', str(text or ''), flags=re.I)
    parts, buffer, depth = [], [], 0
    for char in text:
        if char == '(':
            depth += 1
        elif char == ')':
            depth = max(0, depth - 1)
        if char in ',;\n' and depth == 0:
            parts.append(''.join(buffer).strip())
            buffer = []
        else:
            buffer.append(char)
    parts.append(''.join(buffer).strip())
    result = []
    for part in filter(None, parts):
        # 3 Maja i 11 Listopada to nazwy ulic, a nie kolejne numery.
        if result and re.match(r'^\d+[A-Za-z]?(?:\s*(?:[-–—,/]|\(|$))', part):
            result[-1] += ', ' + part
        else:
            result.append(part)
    return result


def resolve(fragment, index):
    clean = re.sub(r'^(?:ul\.|ulica\s*:)\s*', '', fragment, flags=re.I).strip()
    # Szukamy najdłuższej nazwy; cyfry mogą należeć do nazwy (3 Maja).
    cuts = [len(clean)] + [m.start() for m in re.finditer(r'\s+', clean)][::-1]
    for cut in cuts:
        name, suffix = clean[:cut].strip(), clean[cut:].strip()
        if suffix and not re.match(r'^(?:\d|\(|od\b|do\b)', suffix, re.I):
            continue
        candidates = index.match(name)
        if len(candidates) == 1:
            return candidates[0] + (' ' + suffix if suffix else ''), 'PRG', ''
        if len(candidates) > 1:
            return fragment, 'NIEJEDNOZNACZNE', ' | '.join(candidates)
    suggestions = difflib.get_close_matches(name_key(clean), index.alias.keys(), n=3, cutoff=0.65)
    suggested_names = sorted({n for k in suggestions for n in index.alias[k]})
    return fragment, 'BRAK_W_PRG', ' | '.join(suggested_names)


def process(input_path, gdb, output, report, sheet_name=None):
    import openpyxl
    from openpyxl.styles import Alignment
    if input_path.resolve() == output.resolve():
        raise ValueError('Plik wynikowy musi być inny niż źródłowy.')
    index = StreetIndex(read_prg(gdb))
    workbook = openpyxl.load_workbook(input_path)
    sheets = [workbook[sheet_name]] if sheet_name else list(workbook)
    audit, counts, total, max_parts = [], Counter(), 0, 0
    for sheet in sheets:
        headers = {str(c.value).strip(): c.column for c in sheet[1] if c.value is not None}
        if 'Opis granic' not in headers:
            if sheet_name:
                raise ValueError('Brak kolumny Opis granic w wybranym arkuszu.')
            continue
        if any(re.fullmatch(r'Granica_\d+', h) for h in headers):
            raise ValueError('Arkusz zawiera już kolumny Granica_XX; użyj pliku źródłowego.')
        start = sheet.max_column + 1
        row_results = []
        for row in range(2, sheet.max_row + 1):
            desc = sheet.cell(row, headers['Opis granic']).value
            municipality = sheet.cell(row, headers['Gmina']).value if 'Gmina' in headers else None
            outside = municipality is not None and key(municipality) not in ('m lublin', 'lublin')
            values = []
            for position, fragment in enumerate(split_description(desc), 1):
                value, status, suggestions = ((fragment, 'POZA_LUBLINEM', '') if outside else resolve(fragment, index))
                values.append(value)
                counts[status] += 1
                audit.append([sheet.title, row, municipality, position, fragment, value, status, suggestions])
            row_results.append((row, values))
            max_parts = max(max_parts, len(values))
            total += 1
        width = max((len(v) for _, v in row_results), default=0)
        for offset in range(width):
            cell = sheet.cell(1, start + offset, f'Granica_{offset + 1:02d}')
            cell._style = copy.copy(sheet.cell(1, headers['Opis granic'])._style)
            sheet.column_dimensions[cell.column_letter].width = 45
        for row, values in row_results:
            for offset, value in enumerate(values):
                cell = sheet.cell(row, start + offset, value)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        if sheet.auto_filter.ref:
            sheet.auto_filter.ref = f'A1:{sheet.cell(sheet.max_row, start + width - 1).coordinate}'
    if not total:
        raise ValueError('Nie znaleziono danych w kolumnie Opis granic (nagłówek w wierszu 1).')
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    report.parent.mkdir(parents=True, exist_ok=True)
    with report.open('w', newline='', encoding='utf-8-sig') as stream:
        writer = csv.writer(stream, delimiter=';')
        writer.writerow(['Arkusz', 'Wiersz_Excel', 'Gmina', 'Nr_granicy', 'Fragment_oryginalny', 'Wynik', 'Status', 'Sugestie_do_weryfikacji'])
        writer.writerows(audit)
    print(f'Wiersze: {total}; maksymalna liczba granic: {max_parts}; statusy: {dict(counts)}')
    print(f'Wynik: {output}\nRaport: {report}')
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=BASE / 'obwody_glosowania_Lublin.xlsx')
    parser.add_argument('--gdb', type=Path, default=BASE / 'Adresy_PRG_Geoportal.gdb')
    parser.add_argument('--output', type=Path, default=GEO_RESULTS / '01_granice/obwody_glosowania_Lublin_granice.xlsx')
    parser.add_argument('--report', type=Path, default=GEO_RESULTS / '01_granice/raport_dopasowania_PRG.csv')
    parser.add_argument('--sheet', help='Nazwa arkusza; domyślnie wszystkie z kolumną Opis granic.')
    args = parser.parse_args()
    process(args.input, args.gdb, args.output, args.report, args.sheet)


if __name__ == '__main__':
    main()
